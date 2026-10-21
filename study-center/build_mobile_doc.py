"""Build the mobile DOCX from the learning center content. Use the bundled Python runtime."""
from pathlib import Path
from collections import Counter
import json, re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "study-center/content.json").read_text())
OUT = ROOT / "手机背诵" / "Java后端必备八股文手机背诵版.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)
ORDER = ["Java 基础", "集合", "并发", "JVM", "Spring", "MySQL", "Redis", "网络与 Linux", "分布式与消息", "AI 与工作流"]
LABELS = {"集合":"集合容器", "并发":"Java 并发", "JVM":"JVM 内存与排障", "网络与 Linux":"网络与流式调用", "AI 与工作流":"AI 平台与工作流"}
doc = Document()
section = doc.sections[0]
section.page_width = Mm(148)
section.page_height = Mm(210)
section.top_margin = Mm(12)
section.bottom_margin = Mm(12)
section.left_margin = Mm(12)
section.right_margin = Mm(12)
section.footer_distance = Mm(5)

def font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rp = style.element.get_or_add_rPr()
    fonts = rp.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rp.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Heiti SC")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    for key in ["w:asciiTheme","w:hAnsiTheme","w:eastAsiaTheme","w:cstheme"]:
        fonts.attrib.pop(qn(key), None)
    lang=OxmlElement("w:lang")
    lang.set(qn("w:val"),"en-US")
    lang.set(qn("w:eastAsia"),"zh-CN")
    rp.append(lang)

for s in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    font(doc.styles[s], 11.5)
normal = doc.styles["Normal"]
normal.paragraph_format.line_spacing = Pt(16)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.widow_control = True
font(doc.styles["Title"], 23, True)
doc.styles["Title"].paragraph_format.space_after=Pt(14)
font(doc.styles["Subtitle"], 11.5)
doc.styles["Subtitle"].paragraph_format.space_after=Pt(10)
font(doc.styles["Heading 1"], 17, True)
doc.styles["Heading 1"].paragraph_format.space_before=Pt(17)
doc.styles["Heading 1"].paragraph_format.space_after=Pt(9)
doc.styles["Heading 1"].paragraph_format.keep_with_next=True
font(doc.styles["Heading 2"], 13.5, True)
doc.styles["Heading 2"].paragraph_format.line_spacing=Pt(19)
doc.styles["Heading 2"].paragraph_format.space_before=Pt(12)
doc.styles["Heading 2"].paragraph_format.space_after=Pt(6)
doc.styles["Heading 2"].paragraph_format.keep_with_next=True

# Mobile reading benefits from actual headings, a linked contents page, and page numbers.
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run()
r.font.size=Pt(8)
r.font.color.rgb=RGBColor.from_string("666666")
field=OxmlElement("w:fldSimple")
field.set(qn("w:instr"),"PAGE")
r._r.addnext(field)

def bookmark(p, name, ident):
    a=OxmlElement("w:bookmarkStart");a.set(qn("w:id"),str(ident));a.set(qn("w:name"),name)
    b=OxmlElement("w:bookmarkEnd");b.set(qn("w:id"),str(ident))
    p._p.insert(0,a);p._p.append(b)

def link(p, text, anchor=None, url=None):
    el=OxmlElement("w:hyperlink")
    if anchor:el.set(qn("w:anchor"),anchor)
    if url:
        rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
        el.set(qn("r:id"),rel)
    run=OxmlElement("w:r")
    prop=OxmlElement("w:rPr")
    color=OxmlElement("w:color");color.set(qn("w:val"),"000000");prop.append(color)
    underline=OxmlElement("w:u");underline.set(qn("w:val"),"single");prop.append(underline)
    run.append(prop)
    t=OxmlElement("w:t");t.text=text;run.append(t)
    el.append(run);p._p.append(el)

def body(text, label=None, keep=False):
    p=doc.add_paragraph()
    if label:p.add_run(label+"  ").bold=True
    p.add_run(text)
    p.paragraph_format.keep_with_next=keep
    p.paragraph_format.keep_together=True
    return p

def heading_text(q):
    q=q.replace("==","双等号").replace("_"," ").replace("."," ")
    q=re.sub(r"[？?，,、：:；;／/（）()“”\"‘’'\-—]+"," ",q)
    return re.sub(r"\s+"," ",q).strip()

doc.add_paragraph("Java 后端必备八股文",style="Title")
doc.add_paragraph("手机背诵版  110 题",style="Subtitle")
body("这份手册帮助你先补齐 Java 后端基础，再理解大模型网关、Agent 平台与工作流引擎。每题先读白话理解，再背正式回答，最后用例子和追问检查自己。")
body("每天先学 3 题，之后每次复习 5 题。遮住答案讲 30 至 60 秒；答不准就回看，再用自己的话复述。能背出定义后，还要能举例并说出限制。")
body("复习节奏  当天学习，次日、第 4 天、第 8 天和第 15 天回忆。不要把看懂一次当成记住。")
body("优先级  基础必会先背，项目必会结合代码理解，进阶再学放到基础之后。")
body("版本提示  Java 以 JDK 21 相关概念为主；数据库、中间件和框架的实现细节，应结合项目实际版本与配置核对。")
body("整理日期  2026 年 10 月 21 日")
doc.add_page_break()

p=doc.add_paragraph("章节目录",style="Heading 1")
bookmark(p,"contents",1)
counts=Counter(c["cat"] for c in DATA["cards"])
for n,cat in enumerate(ORDER,1):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(8)
    link(p,f"{n:02d}  {LABELS.get(cat,cat)}  {counts[cat]} 题",anchor="chapter_"+str(n))
body("点击章节可跳转。手机 Word 或 WPS 若支持大纲，也可以通过标题导航定位到具体题目。")
body("每题结构  白话理解是记忆入口；背诵答案保留关键条件；例子帮助联系实际；别说错提醒边界；追问用于闭卷检查。")
doc.add_page_break()

number=0
for chapter,cat in enumerate(ORDER,1):
    p=doc.add_paragraph(f"{chapter:02d}  {LABELS.get(cat,cat)}",style="Heading 1")
    bookmark(p,"chapter_"+str(chapter),10+chapter)
    # Put elementary standalone questions before the integrated project questions.
    cards=sorted([c for c in DATA["cards"] if c["cat"]==cat],key=lambda c:(c["id"].startswith("core-"),int(c["id"].split("-")[1])))
    for c in cards:
        number+=1
        p=doc.add_paragraph(f"{number:03d}  {heading_text(c['q'])}",style="Heading 2")
        bookmark(p,"question_"+str(number),100+number)
        fields=[
            ("学习顺序",c["priority"]),
            ("白话理解",c["plain"]),
            ("背诵答案",c["answer"]),
            ("例子",c.get("example") or "写一个最小例子，验证正常和失败两种情况。"),
        ]
        if c.get("pitfall"):fields.append(("别说错",c["pitfall"]))
        fields.append(("追问",c.get("follow") or "说明它解决什么问题，以及一个常见限制。"))
        for i,(label,text) in enumerate(fields):
            p=body(text,label,keep=i<len(fields)-1)
            if label=="学习顺序":
                p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(9.5)
        # Every question remains a coherent block when it fits on one page.
assert number==110
doc.add_page_break()
doc.add_paragraph("闭卷检查清单",style="Heading 1")
for text in [
    "Java  基本类型与引用、双等号与 equals、异常、泛型和对象共享。",
    "集合  List Set Map、哈希冲突、联合操作原子性、有界队列。",
    "并发  可见性、锁、CAS、线程池、中断、取消和资源释放。",
    "数据库  索引、回表、事务隔离、快照读、锁和提交时机。",
    "中间件  缓存故障、过期与淘汰、消费重复、业务幂等。",
    "AI 平台  SSE 生命周期、总体超时、DAG 汇合、工具调用和 RAG。",
]:
    body(text)
body("每次复习任选两题，按定义、用途、限制、例子的顺序回答。项目实现与个人贡献要依据事实说明，不能用背诵代替源码和运行证据。")
doc.add_paragraph("官方参考",style="Heading 1")
body("需要核对细节时，可打开下面的官方资料。参考版本不代表公司项目当前使用版本。")
for cat,title,url in DATA["sources"]:
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    link(p,title,url=url)
    for r in p.runs:r.font.size=Pt(10)

doc.core_properties.title="Java 后端必备八股文手机背诵版"
doc.core_properties.subject="110 题 Java 后端与 AI 平台基础背诵"
doc.core_properties.author=""
doc.core_properties.keywords="Java 八股文 后端 AI平台 手机背诵"
doc.core_properties.comments=""
for parent in [doc.styles.element, doc.element]:
    for border in list(parent.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.save(OUT)
print(json.dumps({"output":str(OUT),"questions":number,"chapters":len(ORDER),"paragraphs":len(doc.paragraphs)},ensure_ascii=False))

